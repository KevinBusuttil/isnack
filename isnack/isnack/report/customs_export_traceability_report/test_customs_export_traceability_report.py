# Copyright (c) 2026, Busuttil Technologies Limited and contributors
# For license information, please see license.txt

"""Unit tests for the apportioning columns of the Customs Export Traceability
Report (pure mocks, no site).

The report lists the whole consumption of every Work Order that fed a
finished-goods batch. The apportioned figures scale that consumption to the
cartons of the batch actually sold on the invoice line, so their total agrees
with the cost-of-sales posting of the delivery.
"""

import base64
import unittest
from io import BytesIO
from unittest.mock import patch

import frappe

import isnack.isnack.report.customs_export_traceability_report.customs_export_traceability_report as report

FILTERS = frappe._dict(company="Isnack", from_date="2026-08-01", to_date="2026-08-31")


def _si_row(idx, item_code, stock_qty, bundle=None, batch_no=None):
	return frappe._dict(
		sales_invoice="SINV-1",
		company="Isnack",
		posting_date="2026-08-27",
		customer="C",
		customer_name="Customer",
		currency="EUR",
		idx=idx,
		fg_item_code=item_code,
		fg_item_name=item_code.lower(),
		fg_description=item_code,
		qty=stock_qty,
		uom="Carton",
		stock_qty=stock_qty,
		item_group="Finished Goods",
		batch_no=batch_no,
		serial_and_batch_bundle=bundle,
	)


def _wo_entry(work_order, stock_entry, fg_qty, wo_fg_qty=None):
	return {
		"work_order": work_order,
		"stock_entry": stock_entry,
		"manufacturing_date": "2026-08-21",
		"wo_item": "FG10005",
		"wo_qty": fg_qty,
		"fg_qty": fg_qty,
		"wo_fg_qty": wo_fg_qty if wo_fg_qty is not None else fg_qty,
	}


def _rm(item_code, qty, rate, batch_no="RMB-1"):
	return {
		"item_code": item_code,
		"item_name": item_code.lower(),
		"description": item_code,
		"stock_uom": "Kg",
		"qty": qty,
		"consumed_cost": qty * rate,
		"batch_no": batch_no,
		"purchase_receipt": None,
	}


class TestResolveFgBatches(unittest.TestCase):
	def test_direct_batch_takes_the_line_stock_qty(self):
		row = _si_row(1, "FG10005", 449, batch_no="AAO-007")
		self.assertEqual(report._resolve_fg_batches(row, {}), [("AAO-007", 449.0)])

	def test_bundle_splits_the_line_per_batch(self):
		row = _si_row(4, "FG10008", 409, bundle="B1")
		bundles = {"B1": [{"batch_no": "OAA-009", "qty": -405}, {"batch_no": "AAA-004", "qty": -4}]}
		self.assertEqual(
			report._resolve_fg_batches(row, bundles),
			[("OAA-009", 405.0), ("AAA-004", 4.0)],
		)

	def test_delivery_note_bundle_is_scaled_to_the_invoiced_qty(self):
		# The invoice bills 200 of a 449-carton delivery: the DN bundle carries
		# the delivered qty, so the sold qty is scaled to what this line bills.
		row = _si_row(1, "FG10005", 200, bundle="DN-B1")
		bundles = {"DN-B1": [{"batch_no": "AAO-007", "qty": -449}]}
		self.assertEqual(report._resolve_fg_batches(row, bundles), [("AAO-007", 200.0)])

	def test_same_batch_twice_in_a_bundle_is_one_batch(self):
		row = _si_row(1, "FG10005", 10, bundle="B2")
		bundles = {"B2": [{"batch_no": "AAO-007", "qty": -6}, {"batch_no": "AAO-007", "qty": -4}]}
		self.assertEqual(report._resolve_fg_batches(row, bundles), [("AAO-007", 10.0)])

	def test_return_invoice_reverses_the_sold_qty(self):
		# a credit note bills a negative stock qty; the bundle is unsigned, so
		# the sign has to come from the invoice line, as it does for a direct batch
		row = _si_row(1, "FG10005", -449, bundle="RET")
		bundles = {"RET": [{"batch_no": "AAO-007", "qty": 449}]}
		self.assertEqual(report._resolve_fg_batches(row, bundles), [("AAO-007", -449.0)])
		direct = _si_row(1, "FG10005", -449, batch_no="AAO-007")
		self.assertEqual(report._resolve_fg_batches(direct, {}), [("AAO-007", -449.0)])

	def test_no_batch_information(self):
		self.assertEqual(report._resolve_fg_batches(_si_row(1, "FG10005", 10), {}), [(None, None)])
		row = _si_row(1, "FG10005", 10, bundle="EMPTY")
		self.assertEqual(report._resolve_fg_batches(row, {"EMPTY": []}), [(None, None)])


class TestApportionmentFactor(unittest.TestCase):
	def test_sold_share_of_the_batch(self):
		self.assertAlmostEqual(report._apportionment_factor(449, 460, 193, 193), 449 / 460)

	def test_fully_sold_batch_is_one(self):
		self.assertAlmostEqual(report._apportionment_factor(407, 407, 407, 407), 1.0)

	def test_work_order_output_split_over_two_batches(self):
		# 300 of the order's 400 cartons went into this batch, all 300 were sold
		self.assertAlmostEqual(report._apportionment_factor(300, 300, 300, 400), 0.75)

	def test_unknown_quantities_give_none(self):
		self.assertIsNone(report._apportionment_factor(None, 460, 193, 193))
		self.assertIsNone(report._apportionment_factor(449, None, 193, 193))
		self.assertIsNone(report._apportionment_factor(449, 0, 193, 193))

	def test_missing_work_order_output_falls_back_to_the_batch_share(self):
		self.assertAlmostEqual(report._apportionment_factor(449, 460, None, None), 449 / 460)

	def test_apportion_keeps_unknowns_blank(self):
		self.assertIsNone(report._apportion(None, 0.5))
		self.assertIsNone(report._apportion(10, None))
		self.assertAlmostEqual(report._apportion(10, 0.5), 5.0)


class GetDataHarness(unittest.TestCase):
	"""Runs get_data with every database read replaced."""

	def run_report(self, si_items, bundles, wo_map, produced, rm_map):
		with patch.object(report, "_fetch_si_items", return_value=si_items), patch.object(
			report, "_fetch_bundle_entries", return_value=bundles
		), patch.object(report, "_fetch_manufacture_entries", return_value=wo_map), patch.object(
			report, "_fetch_batch_produced_qty", return_value=produced
		), patch.object(report, "_fetch_rm_consumption", return_value=rm_map), patch.object(
			report, "_fetch_pr_details", return_value={}
		), patch.object(report, "_fetch_pr_item_qty", return_value={}), patch.object(
			report, "_fetch_batch_balance", return_value={}
		):
			return report.get_data(FILTERS)


class TestApportionedRows(GetDataHarness):
	def test_batch_made_by_two_work_orders_and_partly_sold(self):
		# AAO-007: 193 + 267 = 460 produced, 449 sold on the line
		rows = self.run_report(
			si_items=[_si_row(1, "FG10005", 449, bundle="B1")],
			bundles={"B1": [{"batch_no": "AAO-007", "qty": -449}]},
			wo_map={("FG10005", "AAO-007"): [_wo_entry("WO-27", "SE-170", 193), _wo_entry("WO-28", "SE-177", 267)]},
			produced={("FG10005", "AAO-007"): 460.0},
			rm_map={
				"WO-27": [_rm("RM20003", 147.91, 13.0), _rm("RM20001", 20, 7.5)],
				"WO-28": [_rm("RM20003", 187, 12.0)],
			},
		)
		self.assertEqual(len(rows), 3)  # one row per raw-material line, as before
		share = 449 / 460
		for r in rows:
			self.assertEqual(r.batch_sold_qty, 449.0)
			self.assertEqual(r.batch_produced_qty, 460.0)
			self.assertAlmostEqual(r.apportioned_qty, r.consumed_qty * share)
			self.assertAlmostEqual(r.apportioned_cost, r.consumed_cost * share)
		# the apportioned total is what the delivery was valued at: sold qty at
		# the batch's average cost per carton
		whole = sum(r.consumed_cost for r in rows)
		self.assertAlmostEqual(sum(r.apportioned_cost for r in rows), 449 * whole / 460)

	def test_line_drawn_from_two_batches(self):
		# 409 cartons: 405 from OAA-009 (fully sold) and 4 from AAA-004 (380 made)
		rows = self.run_report(
			si_items=[_si_row(4, "FG10008", 409, bundle="B4")],
			bundles={"B4": [{"batch_no": "OAA-009", "qty": -405}, {"batch_no": "AAA-004", "qty": -4}]},
			wo_map={
				("FG10008", "OAA-009"): [_wo_entry("WO-13", "SE-103", 405)],
				("FG10008", "AAA-004"): [_wo_entry("WO-04", "SE-030", 380)],
			},
			produced={("FG10008", "OAA-009"): 405.0, ("FG10008", "AAA-004"): 380.0},
			rm_map={"WO-13": [_rm("RM20006", 144, 13.5)], "WO-04": [_rm("RM20006", 128.44, 13.5)]},
		)
		by_batch = {r.fg_batch_no: r for r in rows}
		self.assertEqual(by_batch["OAA-009"].batch_sold_qty, 405.0)
		self.assertAlmostEqual(by_batch["OAA-009"].apportioned_cost, by_batch["OAA-009"].consumed_cost)
		self.assertEqual(by_batch["AAA-004"].batch_sold_qty, 4.0)
		self.assertAlmostEqual(by_batch["AAA-004"].apportioned_qty, 128.44 * 4 / 380)
		self.assertAlmostEqual(by_batch["AAA-004"].apportioned_cost, 128.44 * 13.5 * 4 / 380)
		# the line's own Sales Qty is untouched
		self.assertTrue(all(r.sales_qty == 409 for r in rows))

	def test_work_order_that_booked_into_two_batches(self):
		# WO-9 made 400: 300 into batch X (sold 300) and 100 into batch Y (not on this invoice)
		rows = self.run_report(
			si_items=[_si_row(1, "FG10005", 300, bundle="B1")],
			bundles={"B1": [{"batch_no": "X", "qty": -300}]},
			wo_map={("FG10005", "X"): [_wo_entry("WO-9", "SE-9", 300, wo_fg_qty=400)]},
			produced={("FG10005", "X"): 300.0},
			rm_map={"WO-9": [_rm("RM1", 100, 2.0)]},
		)
		self.assertEqual(len(rows), 1)
		self.assertAlmostEqual(rows[0].apportioned_qty, 75.0)
		self.assertAlmostEqual(rows[0].apportioned_cost, 150.0)

	def test_two_manufacture_entries_into_one_batch_do_not_double_count(self):
		# WO-5 closed twice into the same batch (200 + 100): its whole consumption
		# is listed once, apportioned by both entries' output together.
		rows = self.run_report(
			si_items=[_si_row(1, "FG10005", 300, bundle="B1")],
			bundles={"B1": [{"batch_no": "X", "qty": -300}]},
			wo_map={
				("FG10005", "X"): [
					_wo_entry("WO-5", "SE-A", 200, wo_fg_qty=300),
					_wo_entry("WO-5", "SE-B", 100, wo_fg_qty=300),
				]
			},
			produced={("FG10005", "X"): 300.0},
			rm_map={"WO-5": [_rm("RM1", 90, 1.0)]},
		)
		self.assertEqual(len(rows), 1)
		self.assertAlmostEqual(rows[0].consumed_qty, 90.0)
		self.assertAlmostEqual(rows[0].apportioned_qty, 90.0)
		self.assertEqual(rows[0].manufacture_entry, "SE-A")

	def test_semi_finished_route_reaches_the_row(self):
		rm = dict(_rm("RM20022", 75, 1.0, batch_no="CG1"), via_sfg="SFG1 ← WO-S", sfg_attribution="Single run")
		rows = self.run_report(
			si_items=[_si_row(1, "FG10011", 150, batch_no="BBB-113")],
			bundles={},
			wo_map={("FG10011", "BBB-113"): [_wo_entry("WO-68", "SE-399", 150)]},
			produced={("FG10011", "BBB-113"): 150.0},
			rm_map={"WO-68": [rm, _rm("PM40011", 12, 3.0)]},
		)
		by_item = {r.rm_item_code: r for r in rows}
		self.assertEqual(by_item["RM20022"].via_sfg, "SFG1 ← WO-S")
		self.assertEqual(by_item["RM20022"].sfg_attribution, "Single run")
		self.assertAlmostEqual(by_item["RM20022"].apportioned_qty, 75.0)
		self.assertIsNone(by_item["PM40011"].via_sfg)
		self.assertIsNone(by_item["PM40011"].sfg_attribution)

	def test_batch_without_a_manufacture_entry_stays_blank(self):
		rows = self.run_report(
			si_items=[_si_row(1, "FG10005", 10, bundle="B1")],
			bundles={"B1": [{"batch_no": "BOUGHT-IN", "qty": -10}]},
			wo_map={},
			produced={},
			rm_map={},
		)
		self.assertEqual(len(rows), 1)
		self.assertEqual(rows[0].batch_sold_qty, 10.0)
		self.assertIsNone(rows[0].batch_produced_qty)
		self.assertIsNone(rows[0].apportioned_qty)
		self.assertIsNone(rows[0].apportioned_cost)

	def test_line_without_a_batch_stays_blank(self):
		rows = self.run_report(si_items=[_si_row(1, "FG10005", 10)], bundles={}, wo_map={}, produced={}, rm_map={})
		self.assertEqual(len(rows), 1)
		self.assertIsNone(rows[0].fg_batch_no)
		self.assertIsNone(rows[0].batch_sold_qty)
		self.assertIsNone(rows[0].apportioned_qty)


def _finished_row(sed_name, batch_no, fg_qty, stock_entry="SE-1", work_order="WO-1", item_code="FG10005", scrap=0):
	return frappe._dict(
		stock_entry=stock_entry,
		work_order=work_order,
		manufacturing_date="2026-08-21",
		sed_name=sed_name,
		item_code=item_code,
		batch_no=batch_no,
		is_scrap_item=scrap,
		fg_qty=fg_qty,
	)


def _cons(item_code, qty, rate=1.0, batch_no=None, row=None):
	"""A _consumption_rows dict."""
	return {
		"item_code": item_code,
		"item_name": item_code.lower(),
		"description": item_code,
		"stock_uom": "Kg",
		"qty": qty,
		"consumed_cost": qty * rate,
		"batch_no": batch_no,
		"purchase_receipt": None,
		"row": row or f"row-{item_code}",
		"via_sfg": None,
		"sfg_attribution": None,
	}


def _pool(work_orders=(), other=(), emptied=True):
	return {
		"item_code": "SFG1",
		"warehouse": "Semi-finished - ISN",
		"emptied": emptied,
		"work_orders": [{"work_order": w, "qty": q} for w, q in work_orders],
		"other": list(other),
	}


RECO = {"voucher_type": "Stock Reconciliation", "voucher_no": "MAT-RECO-1", "purpose": None, "posting_date": None, "qty": 20.0}


class TestSemiFinishedDrillDown(unittest.TestCase):
	"""_fetch_rm_consumption for FG WO-68, which drew SFG1 (a mix) and PM1 (film)."""

	def _run(self, consumption, pool, output, made=("SFG1",), filters=None):
		def rows(work_orders):
			return {wo: [dict(r) for r in consumption.get(wo, [])] for wo in work_orders if wo in consumption}

		with patch.object(report, "_consumption_rows", side_effect=rows), patch.object(
			report.batch_lineage, "manufactured_items", side_effect=lambda codes: set(made) & set(codes)
		), patch.object(report.batch_lineage, "pool_sources", side_effect=lambda draws, tick: {
			d: pool[d] for d in draws if d in pool
		}) as sources, patch.object(
			report, "_fetch_wo_finished_qty", side_effect=lambda wos: {w: output[w] for w in wos if w in output}
		), patch.object(report, "qty_tick", return_value=0.001), patch.object(
			report, "_lookup_pr_via_batch", side_effect=lambda batch, item: f"PR-{batch}"
		):
			result = report._fetch_rm_consumption({"WO-68"}, frappe._dict(filters or {}))
		self.sources = sources
		return {(r["item_code"], r.get("batch_no")): r for r in result.get("WO-68", [])}

	FG = {"WO-68": [_cons("SFG1", 80, rate=2.0, row="R-SFG1"), _cons("PM1", 12, batch_no="PM-B")]}

	def test_single_run_replaces_the_mix_by_its_raw_materials(self):
		consumption = dict(self.FG, **{"WO-S": [_cons("RM-CORN", 150, rate=1.0, batch_no="CG1"), _cons("RM-WATER", 10, rate=0.0)]})
		got = self._run(consumption, {"R-SFG1": _pool([("WO-S", 160)])}, {"WO-S": 160.0})

		self.assertNotIn(("SFG1", None), got)
		corn = got[("RM-CORN", "CG1")]
		self.assertAlmostEqual(corn["qty"], 75.0)  # 150 x 80 drawn / 160 made
		self.assertAlmostEqual(corn["consumed_cost"], 75.0)
		self.assertEqual(corn["via_sfg"], "SFG1 ← WO-S")
		self.assertEqual(corn["sfg_attribution"], "Single run")
		self.assertEqual(corn["purchase_receipt"], "PR-CG1")  # customs lookup reaches the grits
		self.assertAlmostEqual(got[("RM-WATER", None)]["qty"], 5.0)  # water is not semi-finished
		film = got[("PM1", "PM-B")]
		self.assertIsNone(film["via_sfg"])
		self.assertEqual(film["qty"], 12)
		self.sources.assert_called_once_with(["R-SFG1"], 0.001)

	def test_several_sources_are_pro_rata_and_labelled(self):
		consumption = dict(
			self.FG,
			**{"WO-C": [_cons("RM-CORN", 60, batch_no="CG1")], "WO-D": [_cons("RM-CORN", 60, batch_no="CG2")]},
		)
		pool = {"R-SFG1": _pool([("WO-C", 60), ("WO-D", 60)], [RECO], emptied=False)}
		got = self._run(consumption, pool, {"WO-C": 60.0, "WO-D": 60.0})

		# 80 drawn over 140 booked: 60/140 from each run, 20/140 of unknown origin
		label = "Pro-rata over 3 sources (estimated)"
		for batch in ("CG1", "CG2"):
			row = got[("RM-CORN", batch)]
			self.assertAlmostEqual(row["qty"], 60 * (80 * 60 / 140) / 60)
			self.assertEqual(row["sfg_attribution"], label)
		self.assertEqual(got[("RM-CORN", "CG1")]["via_sfg"], "SFG1 ← WO-C")
		unknown = got[("SFG1", None)]
		self.assertAlmostEqual(unknown["qty"], 80 * 20 / 140)
		self.assertAlmostEqual(unknown["consumed_cost"], 160 * 20 / 140)
		self.assertEqual(unknown["via_sfg"], "SFG1 ← Stock Reconciliation MAT-RECO-1")
		self.assertEqual(unknown["sfg_attribution"], "Origin not recorded")
		# the parts of the draw add back up to the draw
		self.assertAlmostEqual(sum(r["qty"] for k, r in got.items() if k[0] == "RM-CORN") + unknown["qty"], 80.0)

	def test_no_source_keeps_the_mix_and_says_so(self):
		got = self._run(dict(self.FG), {"R-SFG1": _pool([])}, {})
		mix = got[("SFG1", None)]
		self.assertEqual(mix["qty"], 80)
		self.assertIsNone(mix["via_sfg"])
		self.assertEqual(mix["sfg_attribution"], "Semi-finished, no source found")

	def test_raw_material_filter_applies_inside_the_mix(self):
		consumption = dict(self.FG, **{"WO-S": [_cons("RM-CORN", 150, batch_no="CG1"), _cons("RM-WATER", 10)]})
		got = self._run(
			consumption, {"R-SFG1": _pool([("WO-S", 160)])}, {"WO-S": 160.0}, filters={"raw_material_item": "RM-CORN"}
		)
		self.assertEqual(list(got), [("RM-CORN", "CG1")])

	def test_a_mix_inside_a_mix_chains_the_route_and_keeps_the_weaker_label(self):
		consumption = dict(
			self.FG,
			**{
				"WO-S": [_cons("SFG2", 40, row="R-SFG2"), _cons("RM-CORN", 120, batch_no="CG1")],
				"WO-T": [_cons("RM-OIL", 50, batch_no="OIL1")],
			},
		)
		pool = {
			"R-SFG1": _pool([("WO-S", 160)]),
			"R-SFG2": _pool([("WO-T", 50)], [RECO]),
		}
		got = self._run(consumption, pool, {"WO-S": 160.0, "WO-T": 50.0}, made=("SFG1", "SFG2"))

		oil = got[("RM-OIL", "OIL1")]
		# FG drew half of WO-S; WO-S drew 40 of SFG2, 50/70 from WO-T (made 50)
		self.assertAlmostEqual(oil["qty"], 50 * (40 * 50 / 70) / 50 * 0.5)
		self.assertEqual(oil["via_sfg"], "SFG1 ← WO-S · SFG2 ← WO-T")
		self.assertEqual(oil["sfg_attribution"], "Pro-rata over 2 sources (estimated)")
		unknown = got[("SFG2", None)]
		self.assertEqual(unknown["sfg_attribution"], "Origin not recorded")
		self.assertEqual(unknown["via_sfg"], "SFG1 ← WO-S · SFG2 ← Stock Reconciliation MAT-RECO-1")
		self.assertEqual(got[("RM-CORN", "CG1")]["sfg_attribution"], "Single run")

	def test_depth_cap(self):
		rm_map = {"WO": [_cons("SFG1", 1, row="R")]}
		with patch.object(report.batch_lineage, "manufactured_items") as made:
			self.assertIs(report._expand_semi_finished(rm_map, depth=report.MAX_SFG_DEPTH + 1), rm_map)
		made.assert_not_called()


class TestConsumptionRows(unittest.TestCase):
	def test_quantities_are_in_stock_uom_and_rows_are_named(self):
		raw = [
			frappe._dict(
				work_order="WO-1", sed_name="R1", item_code="RM1", item_name="rm1", description="RM1",
				stock_uom="Kg", qty=25.0, basic_rate=2.0, batch_no="B1", serial_and_batch_bundle=None,
				reference_purchase_receipt="PR-9",
			),
			frappe._dict(
				work_order="WO-1", sed_name="R2", item_code="RM2", item_name="rm2", description="RM2",
				stock_uom="Kg", qty=10.0, basic_rate=1.0, batch_no=None, serial_and_batch_bundle="SBB-1",
				reference_purchase_receipt=None,
			),
		]
		bundles = {"SBB-1": [{"batch_no": "B2", "qty": -6.0}, {"batch_no": "B3", "qty": -4.0}]}
		with patch("frappe.db.sql", return_value=raw) as sql, patch.object(
			report, "_fetch_bundle_entries", return_value=bundles
		):
			rows = report._consumption_rows({"WO-1"})["WO-1"]

		# qty is transfer_qty, the unit stock_uom names and basic_rate prices
		self.assertIn("CASE WHEN IFNULL(sed.transfer_qty, 0) > 0 THEN sed.transfer_qty ELSE sed.qty END AS qty", sql.call_args.args[0])
		self.assertEqual([(r["batch_no"], r["qty"], r["consumed_cost"], r["row"]) for r in rows], [
			("B1", 25.0, 50.0, "R1"), ("B2", 6.0, 6.0, "R2"), ("B3", 4.0, 4.0, "R2"),
		])
		self.assertEqual(rows[0]["purchase_receipt"], "PR-9")
		self.assertTrue(all(r["via_sfg"] is None and r["sfg_attribution"] is None for r in rows))


class TestFinishedQtyPerDetailRow(unittest.TestCase):
	WANTED = {("FG10005", "AAO-007")}

	def test_bundle_rows_of_one_detail_row_are_summed(self):
		# one Serial and Batch Entry per serial number, all in the same batch
		bundle = [_finished_row("row-1", "AAO-007", 1) for _ in range(5)]
		qty = report._finished_qty_per_detail_row([], bundle, self.WANTED)
		self.assertEqual(qty, {("row-1", "FG10005", "AAO-007"): 5.0})

	def test_detail_row_in_both_lookups_is_taken_once(self):
		direct = [_finished_row("row-1", "AAO-007", 193)]
		bundle = [_finished_row("row-1", "AAO-007", 193)]
		qty = report._finished_qty_per_detail_row(direct, bundle, self.WANTED)
		self.assertEqual(qty, {("row-1", "FG10005", "AAO-007"): 193.0})

	def test_other_batches_are_ignored(self):
		bundle = [_finished_row("row-1", "AAO-007", 3), _finished_row("row-1", "OTHER", 2)]
		qty = report._finished_qty_per_detail_row([], bundle, self.WANTED)
		self.assertEqual(qty, {("row-1", "FG10005", "AAO-007"): 3.0})


class TestProducedAndPerEntryQuantities(unittest.TestCase):
	def test_batch_produced_qty_sums_serialised_bundle_rows(self):
		direct = []
		bundle = [_finished_row("row-1", "AAO-007", 1, stock_entry="SE-170") for _ in range(193)] + [
			_finished_row("row-2", "AAO-007", 267, stock_entry="SE-177")
		]
		with patch("frappe.db.sql", side_effect=[direct, bundle]):
			produced = report._fetch_batch_produced_qty({("FG10005", "AAO-007")})
		self.assertEqual(produced, {("FG10005", "AAO-007"): 460.0})

	def test_manufacture_entries_carry_the_summed_entry_qty(self):
		direct = [_finished_row("row-2", "AAO-007", 267, stock_entry="SE-177", work_order="WO-28")]
		bundle = [
			_finished_row("row-1", "AAO-007", 1, stock_entry="SE-170", work_order="WO-27") for _ in range(193)
		] + [_finished_row("row-2", "AAO-007", 267, stock_entry="SE-177", work_order="WO-28")]
		wo_details = [
			frappe._dict(name="WO-27", production_item="FG10005", qty=193, actual_start_date=None),
			frappe._dict(name="WO-28", production_item="FG10005", qty=267, actual_start_date=None),
		]
		with patch("frappe.db.sql", side_effect=[direct, bundle, wo_details]), patch.object(
			report, "_fetch_wo_finished_qty", return_value={"WO-27": 193.0, "WO-28": 267.0}
		):
			wo_map = report._fetch_manufacture_entries({("FG10005", "AAO-007")})
		entries = {e["stock_entry"]: e for e in wo_map[("FG10005", "AAO-007")]}
		self.assertEqual(set(entries), {"SE-177", "SE-170"})
		self.assertEqual(entries["SE-170"]["fg_qty"], 193.0)
		self.assertEqual(entries["SE-177"]["fg_qty"], 267.0)
		self.assertEqual(entries["SE-177"]["wo_fg_qty"], 267.0)

	def test_scrap_rows_book_nothing(self):
		direct = [
			_finished_row("row-1", "AAO-007", 190, stock_entry="SE-170", work_order="WO-27"),
			_finished_row("row-s", "AAO-007", 3, stock_entry="SE-170", work_order="WO-27", scrap=1),
		]
		wo_details = [frappe._dict(name="WO-27", production_item="FG10005", qty=193, actual_start_date=None)]
		with patch("frappe.db.sql", side_effect=[direct, [], wo_details]), patch.object(
			report, "_fetch_wo_finished_qty", return_value={"WO-27": 190.0}
		):
			wo_map = report._fetch_manufacture_entries({("FG10005", "AAO-007")})
		self.assertEqual(wo_map[("FG10005", "AAO-007")][0]["fg_qty"], 190.0)


class TestColumns(unittest.TestCase):
	def test_new_columns_sit_next_to_the_figures_they_qualify(self):
		names = [c["fieldname"] for c in report.get_columns()]
		self.assertEqual(
			names[names.index("fg_batch_no"):names.index("fg_batch_no") + 4],
			["fg_batch_no", "batch_sold_qty", "batch_produced_qty", "work_order"],
		)
		self.assertEqual(
			names[names.index("consumed_qty"):names.index("consumed_qty") + 5],
			["consumed_qty", "consumed_cost", "apportioned_qty", "apportioned_cost", "rm_batch_no"],
		)

	def test_semi_finished_route_follows_the_raw_material(self):
		names = [c["fieldname"] for c in report.get_columns()]
		self.assertEqual(
			names[names.index("rm_description"):names.index("rm_description") + 3],
			["rm_description", "via_sfg", "sfg_attribution"],
		)

	def test_no_existing_column_was_removed(self):
		names = [c["fieldname"] for c in report.get_columns()]
		for expected in (
			"company", "sales_invoice", "posting_date", "customer", "customer_name", "currency",
			"si_item_idx", "fg_item_code", "fg_item_name", "fg_description", "sales_qty", "sales_uom",
			"stock_qty", "fg_net_weight", "fg_gross_weight", "fg_weight_uom", "fg_total_volume",
			"fg_volume_uom", "item_group", "fg_batch_no", "work_order", "wo_item", "wo_qty",
			"manufacturing_date", "manufacture_entry", "rm_item_code", "rm_item_name", "rm_description",
			"rm_uom", "consumed_qty", "consumed_cost", "rm_batch_no", "purchase_receipt",
			"purchase_receipt_date", "supplier_name", "pr_qty", "balance_stock", "customs_document_no",
		):
			self.assertIn(expected, names)


def _report_row(**kw):
	row = frappe._dict(
		company="Isnack",
		sales_invoice="SINV-1",
		posting_date="2026-08-27",
		customer="C",
		customer_name="Customer",
		currency="EUR",
		si_item_idx=1,
		fg_item_code="FG10005",
		fg_item_name="Lentil Squares",
		sales_qty=449.0,
		sales_uom="Carton",
		fg_batch_no="AAO-007",
		batch_sold_qty=449.0,
		batch_produced_qty=460.0,
		work_order="WO-27",
		manufacturing_date="2026-08-21",
		rm_item_code="RM20003",
		rm_item_name="Pellets",
		consumed_qty=147.91,
		consumed_cost=1924.99,
		apportioned_qty=144.373,
		apportioned_cost=1878.958,
		rm_batch_no="512806015",
		purchase_receipt="PR-1",
		purchase_receipt_date="2026-03-19",
		supplier_name="Snack Creations",
		pr_qty=420.0,
		balance_stock=None,
		customs_document_no="SA 651180",
	)
	row.update(kw)
	return row


class TestPrintAndExcelCarryTheColumns(unittest.TestCase):
	def test_print_context_has_the_new_fields(self):
		captured = {}

		def render(template, context):
			captured.update(context)
			return "<html></html>"

		with patch.object(report, "get_data", return_value=[_report_row()]), patch.object(
			report, "_fetch_si_header_details", return_value={}
		), patch("frappe.get_cached_value", return_value="TND"), patch(
			"frappe.render_template", side_effect=render
		):
			report.get_print_html(FILTERS)

		fg = captured["invoices"][0]["fg_items"][0]
		self.assertEqual(fg["batch_sold_qty"], "449.0")
		self.assertEqual(fg["batch_produced_qty"], "460.0")
		row = captured["invoices"][0]["rows"][0]
		self.assertEqual(row["consumed_qty"], "147.91")
		self.assertEqual(row["apportioned_qty"], "144.37")
		self.assertEqual(row["apportioned_cost"], "1878.96")
		self.assertEqual(captured["apportioned_cost_label"], "Apportioned Cost (TND)")

	def test_print_carries_the_semi_finished_route(self):
		captured = {}

		def render(template, context):
			captured.update(context)
			return "<html></html>"

		row = _report_row(via_sfg="SFG1 ← WO-S", sfg_attribution="Single run")
		with patch.object(report, "get_data", return_value=[row]), patch.object(
			report, "_fetch_si_header_details", return_value={}
		), patch("frappe.get_cached_value", return_value="TND"), patch(
			"frappe.render_template", side_effect=render
		):
			report.get_print_html(FILTERS)
		printed = captured["invoices"][0]["rows"][0]
		self.assertEqual(printed["via_sfg"], frappe.utils.escape_html("SFG1 ← WO-S"))
		self.assertEqual(printed["sfg_attribution"], "Single run")

	def test_cost_label_falls_back_without_a_currency(self):
		with patch("frappe.get_cached_value", side_effect=Exception("no site")):
			self.assertEqual(report._apportioned_cost_label("Isnack"), "Apportioned Cost")
		with patch("frappe.get_cached_value", return_value="TND"):
			self.assertEqual(report._apportioned_cost_label("Isnack"), "Apportioned Cost (TND)")

	def test_print_template_lists_the_new_headings(self):
		import os

		path = os.path.join(os.path.dirname(report.__file__), "customs_export_traceability_report_print.html")
		with open(path) as f:
			template = f.read()
		for heading in ("Batch Sold Qty", "Batch Produced Qty", "Apportioned Qty", "{{ apportioned_cost_label }}"):
			self.assertIn(heading, template)
		self.assertIn('<th colspan="6" class="col-group-rm">', template)
		self.assertIn('<div class="via-sfg">via {{ row.via_sfg }} &middot; {{ row.sfg_attribution }}</div>', template)

	def test_excel_export_adds_the_columns_without_dropping_any(self):
		from openpyxl import load_workbook

		row = _report_row(via_sfg="SFG1 ← WO-S", sfg_attribution="Single run")
		with patch.object(report, "get_data", return_value=[row]), patch.object(
			report, "_fetch_si_header_details", return_value={}
		), patch("frappe.get_cached_value", return_value="TND"):
			out = report.get_export_excel(FILTERS)

		ws = load_workbook(BytesIO(base64.b64decode(out["file_content"]))).active
		headers = [[c.value for c in r] for r in ws.iter_rows()]
		fg_header = next(h for h in headers if h[0] == "#")
		rm_header = next(h for h in headers if h[0] == "RM Item Code")
		self.assertEqual(
			fg_header[:15],
			["#", "FG Item Code", "FG Item Name", "Sold Qty", "UOM", "Net Weight", "Gross Weight",
			 "Weight UOM", "Total Volume", "Volume UOM", "FG Batch No", "Batch Sold Qty",
			 "Batch Produced Qty", "Work Order", "Mfg Date"],
		)
		self.assertEqual(
			rm_header[:12],
			["RM Item Code", "RM Item Name", "Consumed Qty", "Apportioned Qty", "Apportioned Cost (TND)",
			 "RM Batch No", "Purchase Receipt", "PR Date", "Supplier Name", "PR Qty", "Balance Stock",
			 "Customs Doc No"],
		)
		fg_row = next(h for h in headers if h[0] == 1)
		self.assertEqual(fg_row[10:13], ["AAO-007", 449.0, 460.0])
		rm_row = next(h for h in headers if h[0] == "RM20003")
		self.assertEqual(rm_row[2], 147.91)
		self.assertAlmostEqual(rm_row[3], 144.373)
		self.assertAlmostEqual(rm_row[4], 1878.958)
		self.assertEqual(rm_row[5], "512806015")
		# the semi-finished route follows the customs columns, under its own group
		self.assertEqual(rm_header[12:14], ["Via Semi-Finished", "SFG Attribution"])
		self.assertEqual(rm_row[12:14], ["SFG1 ← WO-S", "Single run"])
		group = next(h for h in headers if h[0] == "Raw Material Consumed")
		self.assertEqual((group[6], group[12]), ("Purchase / Customs", "Semi-Finished Trace"))


if __name__ == "__main__":
	unittest.main()
